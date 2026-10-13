import copy
import importlib.util
import json
import os
from pathlib import Path
import subprocess
import sys
import tempfile
import unittest

ROOT=Path(__file__).resolve().parents[1]
SCRIPTS=ROOT/'skills/feishu-bitable-browser/scripts'
sys.path.insert(0,str(SCRIPTS))
from workflow import make_plan, source_rows, verify, save


def snapshot():
    return {'url':'https://example.feishu.cn/base/example?table=tblFixture&view=vewFixture',
            'tableId':'tblFixture','viewId':'vewFixture','recordCount':2,
            'fields':[{'id':'t','name':'需求标题','type':1},{'id':'d','name':'需求描述','type':1},
                      {'id':'g','name':'业务模块','type':3,'property':{'options':[{'id':'old','name':'Old','color':1}]}}],
            'records':[{'id':'one','data':{'t':{'value':[{'text':'A'}]},'d':{'value':[{'text':'Description A'}]},'g':{'value':'old'}}},
                       {'id':'two','data':{'t':{'value':[{'text':'B'}]},'d':{'value':[{'text':'Description B'}]},'g':{'value':'old'}}}]}


class MappingTests(unittest.TestCase):
    def setUp(self):
        self.source=[{'row':2,'module':'Group','title':'A','description':'Description A'},
                     {'row':3,'module':'Group','title':'B','description':'Description B'}]

    def test_exact_match_and_noop(self):
        before=snapshot();plan=make_plan(self.source,before,'业务模块','需求标题','需求描述')
        self.assertTrue(plan['ready']);self.assertEqual(plan['change_count'],2)
        before['fields'][2]['property']['options'][0]['name']='Group'
        plan=make_plan(self.source,before,'业务模块','需求标题','需求描述')
        self.assertEqual(plan['change_count'],0);self.assertEqual(plan['new_modules'],[])

    def test_ambiguous_duplicate_blocks_writes(self):
        source=self.source+[dict(self.source[0],row=4)]
        plan=make_plan(source,snapshot(),'业务模块','需求标题','需求描述')
        self.assertFalse(plan['ready']);self.assertTrue(plan['issues'])

    def test_blank_description_not_match_all_empty_rows(self):
        source=[dict(self.source[0],description=''),dict(self.source[1],description='')]
        before=snapshot()
        for r in before['records']:r['data']['d']['value']=None
        self.assertTrue(make_plan(source,before,'业务模块','需求标题','需求描述')['ready'])

    def test_unused_source_requires_explicit_partial(self):
        source=self.source+[{'row':4,'module':'Other','title':'C','description':'C'}]
        self.assertFalse(make_plan(source,snapshot(),'业务模块','需求标题','需求描述')['ready'])
        self.assertTrue(make_plan(source,snapshot(),'业务模块','需求标题','需求描述',True)['ready'])

    def test_incomplete_snapshot_refused(self):
        before=snapshot();before['recordCount']=3
        with self.assertRaises(ValueError):make_plan(self.source,before,'业务模块','需求标题','需求描述')

    def test_other_field_change_detected(self):
        before=snapshot();plan=make_plan(self.source,before,'业务模块','需求标题','需求描述');after=copy.deepcopy(before)
        after['fields'][2]['property']['options'].append({'id':'new','name':'Group','color':2})
        for r in after['records']:r['data']['g']={'value':'new'}
        self.assertTrue(verify(plan,before,after,set())['passed'])
        after['records'][0]['data']['t']['value'][0]['text']='Unexpected title'
        self.assertFalse(verify(plan,before,after,set())['passed'])

    def test_merged_group_not_unbounded_forward_fill(self):
        from openpyxl import Workbook
        with tempfile.TemporaryDirectory() as directory:
            path=Path(directory)/'synthetic.xlsx';w=Workbook();s=w.active;s.title='Source'
            s.append(['Group','Title','Description']);s.append(['Module\nGoal note','A','Desc A']);s.append([None,'B','Desc B'])
            s.merge_cells('A2:A3');w.save(path);w.close()
            rows=source_rows(path,'Source','A','B','C',1,'first-line')
            self.assertEqual([r['module'] for r in rows],['Module','Module'])
            w=Workbook();s=w.active;s.title='Source';s.append(['Group','Title','Description']);s.append(['Module','A','Desc A']);s.append([None,'B','Desc B']);w.save(path);w.close()
            with self.assertRaises(ValueError):source_rows(path,'Source','A','B','C',1,'full-cell')

    def test_backup_cannot_be_overwritten(self):
        with tempfile.TemporaryDirectory() as directory:
            path=Path(directory)/'before.json';save(path,{'original':True})
            with self.assertRaises(FileExistsError):save(path,{'original':False})
            self.assertTrue(json.loads(path.read_text())['original'])

    def test_installer_collision_and_backup(self):
        with tempfile.TemporaryDirectory() as directory:
            dest=Path(directory)/'skills';env=dict(os.environ,FEISHU_SKILL_RUNTIME=str(Path(directory)/'runtime'))
            command=[sys.executable,str(ROOT/'install_skill.py'),'--destination',str(dest)]
            subprocess.run(command,env=env,check=True,capture_output=True)
            sentinel=dest/'feishu-bitable-browser'/'sentinel';sentinel.write_text('preserve')
            self.assertNotEqual(subprocess.run(command,env=env,capture_output=True).returncode,0)
            self.assertTrue(sentinel.exists())
            subprocess.run([*command,'--replace'],env=env,check=True,capture_output=True)
            self.assertTrue(list((Path(directory)/'runtime'/'skill-backups').glob('*/sentinel')))


if __name__=='__main__':unittest.main()
