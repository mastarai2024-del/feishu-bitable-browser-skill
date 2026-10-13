"""Deterministic, offline matching and verification. Never writes to Feishu."""
import argparse
from collections import Counter
import hashlib
import json
import os
from pathlib import Path
import re


def load(path):
    return json.loads(Path(path).read_text(encoding="utf-8"))


def save(path, data):
    path = Path(path).expanduser()
    path.parent.mkdir(parents=True, exist_ok=True)
    # Refuse overwrite, including a previous baseline or receipt.
    fd = os.open(path, os.O_WRONLY | os.O_CREAT | os.O_EXCL, 0o600)
    with os.fdopen(fd, "w", encoding="utf-8") as out:
        json.dump(data, out, ensure_ascii=False, indent=2)
        out.write("\n")


def fingerprint(data):
    text = json.dumps(data, ensure_ascii=False, sort_keys=True, separators=(",", ":"))
    return hashlib.sha256(text.encode()).hexdigest()


def normalized(value):
    return re.sub(r"\s+", "", str(value or ""))


def text(cell):
    value = (cell or {}).get("value")
    if value is None:
        return ""
    if isinstance(value, list):
        return "".join(v.get("text", "") if isinstance(v, dict) else str(v) for v in value)
    return str(value)


def field(snapshot, name):
    matches = [f for f in snapshot["fields"] if f["name"] == name]
    if len(matches) != 1:
        raise ValueError(f"Field must be uniquely identified: {name}")
    return matches[0]


def source_rows(path, sheet_name, group_column, title_column, description_column, header_row, group_label):
    from openpyxl import load_workbook
    from openpyxl.utils.cell import column_index_from_string
    book = load_workbook(path, data_only=True)
    try:
        sheet = book[sheet_name]
        if sheet.max_row > 10000 or sheet.max_column > 100:
            raise ValueError("Source exceeds bounded extraction; export a smaller task range.")
        columns = [column_index_from_string(c) for c in (group_column, title_column, description_column)]
        merge_origin = {}
        for span in sheet.merged_cells.ranges:
            for col in columns:
                if span.min_col <= col <= span.max_col:
                    for row in range(span.min_row, span.max_row + 1):
                        merge_origin[row, col] = (span.min_row, span.min_col)

        def value(row, col):
            row, col = merge_origin.get((row, col), (row, col))
            return sheet.cell(row, col).value

        result = []
        for row in range(header_row + 1, sheet.max_row + 1):
            # A merged secondary title must not invent an otherwise empty requirement.
            raw_title = sheet.cell(row, columns[1]).value
            description = sheet.cell(row, columns[2]).value
            if not raw_title and not description:
                continue
            group = value(row, columns[0])
            if not group:
                raise ValueError(f"Row {row}: missing group outside a merged range")
            module = str(group).strip()
            if group_label == "first-line":
                module = module.splitlines()[0].strip()
            result.append({"row": row, "module": module, "group_source": str(group),
                           "title": str(raw_title or description or ""), "description": str(description or "")})
        if not result:
            raise ValueError("Source has no requirements in the selected columns")
        return result
    finally:
        book.close()


def make_plan(source, snapshot, group_name, title_name, description_name, allow_partial=False):
    if snapshot["recordCount"] != len(snapshot["records"]):
        raise ValueError("Target snapshot is incomplete")
    ids = [r["id"] for r in snapshot["records"]]
    if len(set(ids)) != len(ids):
        raise ValueError("Duplicate target record IDs")
    group = field(snapshot, group_name)
    if group["type"] != 3:
        raise ValueError("Automated mapping supports existing single-select fields only; use UI for other types")
    title = field(snapshot, title_name)
    description = field(snapshot, description_name)
    options = {o["id"]: o["name"] for o in group.get("property", {}).get("options", [])}
    if len(set(options.values())) != len(options):
        raise ValueError("Duplicate option names require UI reconciliation")
    records, issues, used = [], [], set()
    for record in snapshot["records"]:
        data = record["data"]
        target_title, target_desc = text(data.get(title["id"])), text(data.get(description["id"]))
        candidates = [s for s in source if normalized(s["title"]) == normalized(target_title)
                      and normalized(s["description"]) == normalized(target_desc)]
        if not candidates and normalized(target_desc):
            candidates = [s for s in source if normalized(s["description"]) == normalized(target_desc)]
        if not candidates and normalized(target_title):
            candidates = [s for s in source if normalized(s["title"]) == normalized(target_title)]
        if len(candidates) != 1 or candidates[0]["row"] in used:
            issues.append({"record_id": record["id"], "title": target_title,
                           "reason": "missing, ambiguous, or source already matched", "candidate_rows": [s["row"] for s in candidates]})
            continue
        match = candidates[0]
        used.add(match["row"])
        old_value = data.get(group["id"], {}).get("value")
        records.append({"record_id": record["id"], "source_row": match["row"], "module": match["module"],
                        "old_value": old_value, "changed": options.get(old_value) != match["module"],
                        "expected_title_cell": data.get(title["id"]), "expected_description_cell": data.get(description["id"])})
    unused = [s["row"] for s in source if s["row"] not in used]
    modules = list(dict.fromkeys(s["module"] for s in source))
    return {"format": 1, "target_url": snapshot.get("url"), "table_id": snapshot["tableId"],
            "view_id": snapshot.get("viewId"), "field_id": group["id"], "title_field_id": title["id"],
            "description_field_id": description["id"], "record_count": snapshot["recordCount"],
            "baseline_field": group, "baseline_field_sha256": fingerprint(group), "modules": modules,
            "new_modules": [m for m in modules if m not in options.values()], "records": records,
            "issues": issues, "unused_source_rows": unused, "allow_partial": allow_partial,
            "ready": bool(records) and (allow_partial or (not issues and not unused)),
            "change_count": sum(r["changed"] for r in records)}


def verify(plan, before, after, allowed_system_fields):
    old = {r["id"]: r["data"] for r in before["records"]}
    new = {r["id"]: r["data"] for r in after["records"]}
    if set(old) != set(new) or len(new) != plan["record_count"]:
        raise ValueError("Record identities/count changed")
    after_fields = {f["id"]: f for f in after["fields"]}
    choices = after_fields[plan["field_id"]]["property"]["options"]
    labels = {o["id"]: o["name"] for o in choices}
    mismatches = [r["record_id"] for r in plan["records"]
                  if labels.get(new[r["record_id"]].get(plan["field_id"], {}).get("value")) != r["module"]]
    intended = {r["record_id"] for r in plan["records"]}
    unrelated = []
    for rid in old:
        for fid in set(old[rid]) | set(new[rid]):
            if fid in allowed_system_fields or (rid in intended and fid == plan["field_id"]):
                continue
            if old[rid].get(fid) != new[rid].get(fid):
                unrelated.append({"record_id": rid, "field_id": fid})
    old_options = plan["baseline_field"].get("property", {}).get("options", [])
    options_preserved = all(next((n for n in choices if n["id"] == o["id"]), None) == o for o in old_options)
    schema_changes = [f["id"] for f in before["fields"]
                      if f["id"] != plan["field_id"] and after_fields.get(f["id"]) != f]
    target_schema = after_fields[plan["field_id"]]
    expected_schema = dict(plan["baseline_field"])
    expected_schema["property"] = dict(expected_schema.get("property", {}))
    expected_schema["property"]["options"] = choices
    target_schema_preserved = target_schema == expected_schema
    return {"passed": not mismatches and not unrelated and not schema_changes and options_preserved and target_schema_preserved,
            "matched": len(plan["records"]) - len(mismatches), "mismatches": mismatches,
            "unrelated_cell_changes": unrelated, "unrelated_schema_changes": schema_changes,
            "old_options_preserved": options_preserved, "target_schema_preserved": target_schema_preserved,
            "distribution": dict(Counter(labels.get(new[r["record_id"]].get(plan["field_id"], {}).get("value")) for r in plan["records"]))}


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    commands = parser.add_subparsers(dest="command", required=True)
    p = commands.add_parser("map")
    p.add_argument("--source-xlsx", required=True); p.add_argument("--sheet", required=True)
    p.add_argument("--snapshot", required=True); p.add_argument("--output", required=True)
    p.add_argument("--group-column", default="B"); p.add_argument("--title-column", default="C")
    p.add_argument("--description-column", default="D"); p.add_argument("--header-row", type=int, default=1)
    p.add_argument("--group-label", choices=["full-cell", "first-line"], default="full-cell")
    p.add_argument("--field", default="业务模块"); p.add_argument("--title-field", default="需求标题")
    p.add_argument("--description-field", default="需求描述"); p.add_argument("--allow-partial", action="store_true")
    p = commands.add_parser("verify")
    p.add_argument("--plan", required=True); p.add_argument("--before", required=True)
    p.add_argument("--after", required=True); p.add_argument("--output", required=True)
    p.add_argument("--allow-system-field", action="append", default=[])
    args = parser.parse_args()
    if args.command == "map":
        source = source_rows(args.source_xlsx, args.sheet, args.group_column, args.title_column,
                             args.description_column, args.header_row, args.group_label)
        result = make_plan(source, load(args.snapshot), args.field, args.title_field, args.description_field, args.allow_partial)
        save(args.output, result)
        print(json.dumps({k: result[k] for k in ["ready", "change_count", "new_modules", "issues", "unused_source_rows"]}, ensure_ascii=False))
        if not result["ready"]: raise SystemExit(2)
    else:
        result = verify(load(args.plan), load(args.before), load(args.after), set(args.allow_system_field))
        save(args.output, result)
        print(json.dumps(result, ensure_ascii=False))
        if not result["passed"]: raise SystemExit(3)


if __name__ == "__main__":
    try:
        main()
    except (ValueError, KeyError, OSError) as error:
        raise SystemExit(str(error))
